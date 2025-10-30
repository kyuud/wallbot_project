import customtkinter as ctk
from tkinter import filedialog, messagebox
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service as ChromeService
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (ElementClickInterceptedException, NoSuchElementException, TimeoutException, StaleElementReferenceException)
from openpyxl import load_workbook, Workbook
import time
import threading
from datetime import datetime

executando_macro = True

# Função para salvar o progresso
def salvar_progresso(progresso):
    with open("progresso.txt", "w") as file:
        file.write(str(progresso))

# Função para carregar o progresso
def carregar_progresso():
    try:
        with open("progresso.txt", "r") as file:
            return int(file.read().strip())
    except FileNotFoundError:
        return 0

def fechamento_em_lote():
    global executando_macro
    caminho_arquivo = entrada_arquivo.get()
    script = entrada_script.get("1.0", "end-1c")

    # Abrir arquivo de log em modo append
    log_file = open('log_output.txt', 'a', encoding='utf-8')

    # Função para escrever no log com data e hora
    def escrever_log(mensagem):
        agora = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        log_file.write(f'[{agora}] {mensagem}\n')

    # Lista para armazenar os resultados
    results = []

    if executando_macro:
        mensagem_inicio = "Executando macro..."
        print(mensagem_inicio)
        escrever_log(mensagem_inicio)

    chrome_options = Options()
    chrome_options = webdriver.ChromeOptions()
    chrome_options.add_argument('--ignore-certificate-errors')
    chrome_options.add_argument('--ignore-ssl-errors')
    chrome_options.add_argument('--start-maximized')

    # Acessando o SIPCS
    while True:
        try:
            driver = webdriver.Chrome(service=ChromeService(ChromeDriverManager().install()), options=chrome_options)
            url = "https://cartoes.extracaixa/"
            driver.get(url)
            sipcs = driver.current_window_handle

            # Logando no SIPCS
            usuario = driver.find_element(By.NAME, 'loginForm:username')
            senha = driver.find_element(By.NAME, 'loginForm:password')
            # usuario.send_keys(email)
            # senha.send_keys(password)
            usuario.send_keys('wbsouza1@stefanini.com')
            senha.send_keys('Barbara05')
            botao_login = driver.find_element(By.XPATH, '//*[@id="loginForm"]/div/div/input')
            botao_login.click()

            # Acessando o SIACH
            opcao_siach = WebDriverWait(driver, 10).until(EC.visibility_of_element_located((By.XPATH, '//*[@id="index"]/fieldset/div[1]/table/tbody/tr[1]/td[2]/a')))  
            opcao_siach.click()
            break

        except Exception as e:
            driver.quit()
            mensagem_erro = "Erro de comunicação. Tentando novamente..."
            print(mensagem_erro)
            escrever_log(mensagem_erro)

    # Função para ler os valores da coluna A da planilha Excel
    def ler_valores_coluna_a(caminho_arquivo, nome_aba):
        workbook = load_workbook(filename=caminho_arquivo)
        sheet = workbook[nome_aba]
        valores = []
        for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0] is None:
                break
            valores.append(row[0])
        return valores

    caminho_arquivo_excel = caminho_arquivo
    nome_aba = 'Planilha1'
    valores_coluna_a = ler_valores_coluna_a(caminho_arquivo_excel, nome_aba)

    # Alterar a aba para SIACH
    abas_handles = driver.window_handles
    siach = None
    for handle in abas_handles:
        if handle != sipcs:
            siach = handle
            break
    driver.switch_to.window(siach)
    driver.execute_script("document.body.style.zoom='0.67'")
    # Carregar o progresso salvo
    progresso = carregar_progresso()
    total_protocolos = len(valores_coluna_a)
    progress_bar.set(progresso / total_protocolos)

    # Loop para tratar protocolo
    protocolos_processados = 0
    inicio = time.time()
    try:
        for i in range(progresso, len(valores_coluna_a)):
            if not executando_macro:
                driver.quit()
                mensagem_interrupcao = "Macro interrompida pelo usuário."
                atualizar_frame_resultado(mensagem_interrupcao)
                escrever_log(mensagem_interrupcao)
                return

            valor = valores_coluna_a[i]
            try:
                opcao_atender_ocorrencia = WebDriverWait(driver, 10).until(EC.visibility_of_element_located((By.XPATH, '//*[@id="menu"]/li[3]/a'))) 
                time.sleep(0.5)
                driver.execute_script("arguments[0].click();", opcao_atender_ocorrencia)
                time.sleep(2)
                campo_texto = driver.find_element(By.XPATH, '/html/body/div[2]/div/div/div[2]/div[1]/form/div[1]/div[1]/div[2]/div/input')
                campo_texto.clear()
                campo_texto.send_keys(valor)
                botao_consultar = driver.find_element(By.XPATH, '/html/body/div[2]/div/div/div[2]/div[1]/form/div[2]/a[1]')
                time.sleep(0.5)
                driver.execute_script("arguments[0].click();", botao_consultar)
                fase = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.XPATH, '/html/body/div[2]/div/div/div[2]/div[2]/div/div[2]/div/div/div/table/tbody/tr/td/div/div[1]/div/div[2]/p[3]/span')))

                try:
                    fase_texto = fase.text
                except StaleElementReferenceException:
                    fase = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.XPATH, '/html/body/div[2]/div/div/div[2]/div[2]/div/div[2]/div/div/div/table/tbody/tr/td/div/div[1]/div/div[2]/p[3]/span')))
                    fase_texto = fase.text

                if fase_texto in ["ABERTA", "EM ANDAMENTO"]:
                    botao_lupa = driver.find_element(By.XPATH, '//*[@id="detalhe_ocorrencia"]/span')
                    time.sleep(0.5)
                    driver.execute_script("arguments[0].click();", botao_lupa)
                    time.sleep(2)
                    try:
                        botao_finalizar = WebDriverWait(driver, 10).until(EC.element_to_be_clickable((By.XPATH, '//*[@id="content"]/div/div[6]/div/div[3]/form/a[6]')))
                        time.sleep(0.5)
                        driver.execute_script("arguments[0].click();", botao_finalizar)
                        time.sleep(2)
                        campo_justificativa = driver.find_element(By.XPATH, '/html/body/div[2]/div/div/div[2]/div[1]/form/div[1]/fieldset[1]/div/div[2]/div/input')
                        campo_justificativa.send_keys('FIM')
                        campo_respostacliente = driver.find_element(By.XPATH, '/html/body/div[2]/div/div/div[2]/div[1]/form/div[1]/fieldset[1]/div/div[4]/div/textarea')
                        campo_respostacliente.send_keys(script)
                        campo_informacaotecnica = driver.find_element(By.XPATH, '/html/body/div[2]/div/div/div[2]/div[1]/form/div[1]/fieldset[1]/div/div[7]/div/textarea')
                        campo_informacaotecnica.send_keys(script)

                        time.sleep(3)

                        try:
                            botao_salvar = WebDriverWait(driver, 20).until(
                                EC.element_to_be_clickable((By.XPATH, "//a[@data-ng-click='tratarOcorrenciaCtrl.form.submit()']")))
                            driver.execute_script("arguments[0].click();", botao_salvar)
                        except Exception as e:
                            mensagem = f"Ocorreu um erro ao clicar no botão Salvar: {e}"
                            print(mensagem)
                            escrever_log(mensagem)
                            driver.save_screenshot(f'erro_ao_clicar_salvar_{valor}.png')
                            # Registrar resultado
                            results.append({
                                'Protocolo': valor,
                                'Status': 'Erro',
                                'Mensagem': mensagem
                            })
                            continue

                        time.sleep(3)

                        try:
                            botao_sim = WebDriverWait(driver, 20).until(EC.element_to_be_clickable((By.CSS_SELECTOR, "button.btn.btn-default")))
                            driver.execute_script("arguments[0].click();", botao_sim)
                            time.sleep(2)
                        except Exception as e:
                            mensagem = f"Ocorreu um erro ao clicar no botão Sim: {e}"
                            print(mensagem)
                            escrever_log(mensagem)
                            driver.save_screenshot(f'erro_ao_clicar_sim_{valor}.png')
                            results.append({
                                'Protocolo': valor,
                                'Status': 'Erro',
                                'Mensagem': mensagem
                            })
                            continue

                        mensagem = f'Concluído: Protocolo {valor}'
                        print(mensagem)
                        atualizar_frame_resultado(mensagem)
                        escrever_log(mensagem)
                        results.append({
                            'Protocolo': valor,
                            'Status': 'Concluído',
                            'Mensagem': mensagem
                        })

                    except (TimeoutException, NoSuchElementException):
                        mensagem = f'Protocolo {valor} já finalizado, mesmo tendo fase {fase_texto}'
                        print(mensagem)
                        atualizar_frame_resultado(mensagem)
                        escrever_log(mensagem)
                        results.append({
                            'Protocolo': valor,
                            'Status': 'Já finalizado',
                            'Mensagem': mensagem
                        })

                elif fase_texto in ["REABERTA", "FINALIZADA"]:
                    mensagem = f'Sem ação, Protocolo {valor} com fase {fase_texto}'
                    print(mensagem)
                    atualizar_frame_resultado(mensagem)
                    escrever_log(mensagem)
                    results.append({
                        'Protocolo': valor,
                        'Status': 'Sem ação',
                        'Mensagem': mensagem
                    })

                else:
                    mensagem = f'Fase desconhecida: Protocolo {valor} com fase {fase_texto}'
                    print(mensagem)
                    atualizar_frame_resultado(mensagem)
                    escrever_log(mensagem)
                    # Registrar resultado
                    results.append({
                        'Protocolo': valor,
                        'Status': 'Fase desconhecida',
                        'Mensagem': mensagem
                    })

                protocolos_processados += 1

                # Salvar progresso
                salvar_progresso(i + 1)

                # Atualizar barra de progresso
                progress_bar.set((i + 1) / total_protocolos)
                progress_label.configure(text=f"Progresso: {i + 1}/{total_protocolos}")

            except ElementClickInterceptedException as e:
                mensagem = f'ElementClickInterceptedException: Protocolo {valor}'
                print(mensagem)
                atualizar_frame_resultado(mensagem)
                escrever_log(mensagem)
                # Registrar resultado
                results.append({
                    'Protocolo': valor,
                    'Status': 'Erro',
                    'Mensagem': mensagem
                })

            except TimeoutException as e:
                mensagem = f'TimeoutException: Protocolo {valor}'
                print(mensagem)
                atualizar_frame_resultado(mensagem)
                escrever_log(mensagem)
                driver.quit()
                # Registrar resultado
                results.append({
                    'Protocolo': valor,
                    'Status': 'Erro',
                    'Mensagem': mensagem
                })
                fechamento_em_lote()
                return

    finally:
        driver.quit()
        fim = time.time()
        tempo_final = fim - inicio
        mensagem_final = f'O lote foi finalizado em {tempo_final:.2f} segundos'
        atualizar_frame_resultado(mensagem_final)
        print(mensagem_final)
        escrever_log(mensagem_final)
        log_file.close()

        # Escrever resultados em um arquivo Excel
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Resultados"

        # Escrever cabeçalhos
        sheet.append(['Protocolo', 'Status', 'Mensagem'])

        # Escrever dados
        for result in results:
            sheet.append([result['Protocolo'], result['Status'], result['Mensagem']])

        # Salvar o arquivo Excel
        workbook.save('resultado_saida.xlsx')

        messagebox.showinfo("Processamento concluído", f"Processamento concluído, {protocolos_processados} protocolos processados")

def parar_macro():
    global executando_macro
    executando_macro = False
    print("Macro parada.")

def atualizar_frame_resultado(texto):
    text_resultado.insert("end", texto + "\n")
    text_resultado.see("end")

def iniciar_fechamento_em_lote():
    global executando_macro
    executando_macro = True
    threading.Thread(target=fechamento_em_lote).start()

# Configurar o estilo do CustomTkinter
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

root = ctk.CTk()
root.title("WallBot - Finalização de Ocorrências SIACH v1.1")
root.geometry("600x500")

# Frame superior
frame_top = ctk.CTkFrame(root)
frame_top.pack(padx=20, pady=10, fill="x")

# Label, Entry e Botão para o arquivo Excel
ctk.CTkLabel(frame_top, text="Caminho do Arquivo:").grid(row=2, column=0, sticky="w")
entrada_arquivo = ctk.CTkEntry(frame_top, width=250)
entrada_arquivo.grid(row=2, column=1, padx=10, pady=5)
ctk.CTkButton(
    frame_top,
    text="Selecionar Arquivo",
    command=lambda: entrada_arquivo.insert("end", filedialog.askopenfilename())
).grid(row=2, column=2, padx=10, pady=5)

# Label, Entry e Botão para o script
ctk.CTkLabel(frame_top, text="Script de Finalização:").grid(row=3, column=0, sticky="nw")
entrada_script = ctk.CTkTextbox(frame_top, width=250, height=100)
entrada_script.grid(row=3, column=1, padx=10, pady=5)
ctk.CTkButton(
    frame_top,
    text="Limpar",
    command=lambda: entrada_script.delete("1.0", "end")
).grid(row=3, column=2, padx=10, pady=5, sticky="n")

# Frame para os botões
frame_botoes = ctk.CTkFrame(root)
frame_botoes.pack(padx=20, pady=10, fill="x")

# Botão de iniciar e parar a execução
ctk.CTkButton(
    frame_botoes,
    text="Iniciar",
    command=iniciar_fechamento_em_lote
).pack(side="left", padx=10, pady=10)
ctk.CTkButton(
    frame_botoes,
    text="Parar",
    command=parar_macro
).pack(side="left", padx=10, pady=10)

# Barra de progresso e rótulo
frame_progresso = ctk.CTkFrame(root)
frame_progresso.pack(padx=20, pady=10, fill="x")

progress_label = ctk.CTkLabel(frame_progresso, text="Progresso: 0/0")
progress_label.pack(anchor="w")

progress_bar = ctk.CTkProgressBar(frame_progresso, width=400)
progress_bar.pack(pady=5, fill="x")
progress_bar.set(0)

# Text para exibir resultado com Scrollbar
frame_resultado = ctk.CTkFrame(root)
frame_resultado.pack(padx=20, pady=10, fill="both", expand=True)

text_resultado = ctk.CTkTextbox(frame_resultado, wrap="word")
text_resultado.pack(side="left", fill="both", expand=True, padx=10, pady=5)

scrollbar = ctk.CTkScrollbar(frame_resultado, command=text_resultado.yview)
scrollbar.pack(side="right", fill="y")
text_resultado.configure(yscrollcommand=scrollbar.set)

root.mainloop()
