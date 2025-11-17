import customtkinter as ctk
from tkinter import filedialog, messagebox
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.firefox.service import Service as FirefoxService
from selenium.webdriver.firefox.options import Options
from webdriver_manager.firefox import GeckoDriverManager
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    ElementClickInterceptedException, 
    NoSuchElementException, 
    TimeoutException, 
    StaleElementReferenceException,
    WebDriverException
)
from openpyxl import load_workbook, Workbook
import time
import threading
from datetime import datetime
import os

# ==================== CONFIGURAÇÕES ====================
CONFIG = {
    'URL_SIPCS': 'https://cartoes.extracaixa/',
    'ARQUIVO_PROGRESSO': 'progresso.txt',
    'ARQUIVO_LOG': 'log_output.txt',
    'ARQUIVO_RESULTADO': 'resultado_saida.xlsx',
    'NOME_ABA_EXCEL': 'Planilha1',
    'ZOOM_NIVEL': '0.67',
    'TIMEOUT_PADRAO': 15,
    'TIMEOUT_SALVAR': 30,
    'MAX_TENTATIVAS_RECONEXAO': 3,
    'ESPERA_ENTRE_ACOES': 1.0
}

# ==================== GERENCIAMENTO DE ESTADO ====================
class EstadoAplicacao:
    def __init__(self):
        self._executando = True
        self._lock = threading.Lock()
    
    def esta_executando(self):
        with self._lock:
            return self._executando
    
    def parar(self):
        with self._lock:
            self._executando = False
    
    def iniciar(self):
        with self._lock:
            self._executando = True

class Estatisticas:
    def __init__(self):
        self.concluidos = 0
        self.sem_acao = 0
        self.erros = 0
        self.fase_desconhecida = 0
        self._lock = threading.Lock()
    
    def adicionar(self, status):
        with self._lock:
            if status == 'Concluído':
                self.concluidos += 1
            elif status in ['Sem ação', 'Já finalizado']:
                self.sem_acao += 1
            elif status == 'Fase desconhecida':
                self.fase_desconhecida += 1
            else:  # Erro
                self.erros += 1
    
    def obter_totais(self):
        with self._lock:
            return {
                'concluidos': self.concluidos,
                'sem_acao': self.sem_acao,
                'erros': self.erros,
                'fase_desconhecida': self.fase_desconhecida,
                'total': self.concluidos + self.sem_acao + self.erros + self.fase_desconhecida
            }
    
    def resetar(self):
        with self._lock:
            self.concluidos = 0
            self.sem_acao = 0
            self.erros = 0
            self.fase_desconhecida = 0

estado_app = EstadoAplicacao()
estatisticas = Estatisticas()

# ==================== GERENCIAMENTO DE LOG ====================
class GerenciadorLog:
    def __init__(self, arquivo):
        self.arquivo = arquivo
        self.file_handle = None
    
    def __enter__(self):
        self.file_handle = open(self.arquivo, 'a', encoding='utf-8')
        return self
    
    def __exit__(self, exc_type, exc_val, exc_tb):
        if self.file_handle:
            self.file_handle.close()
    
    def escrever(self, mensagem):
        agora = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        linha = f'[{agora}] {mensagem}\n'
        if self.file_handle:
            self.file_handle.write(linha)
            self.file_handle.flush()
        print(mensagem)

# ==================== GERENCIAMENTO DE PROGRESSO ====================
def salvar_progresso(progresso):
    """Salva o progresso atual"""
    try:
        with open(CONFIG['ARQUIVO_PROGRESSO'], "w") as file:
            file.write(str(progresso))
    except Exception as e:
        print(f"Erro ao salvar progresso: {e}")

def carregar_progresso():
    """Carrega o progresso salvo"""
    try:
        with open(CONFIG['ARQUIVO_PROGRESSO'], "r") as file:
            return int(file.read().strip())
    except FileNotFoundError:
        return 0
    except ValueError:
        print("Arquivo de progresso corrompido. Iniciando do zero.")
        return 0

def resetar_progresso():
    """Reseta o progresso para zero"""
    salvar_progresso(0)

# ==================== GERENCIAMENTO DE EXCEL ====================
def ler_valores_coluna_a(caminho_arquivo, nome_aba):
    """Lê os valores da coluna A da planilha Excel"""
    try:
        # Validar se arquivo existe
        if not os.path.exists(caminho_arquivo):
            raise Exception(f"Arquivo não encontrado: {caminho_arquivo}")
        
        workbook = load_workbook(filename=caminho_arquivo)
        
        # Validar se aba existe
        if nome_aba not in workbook.sheetnames:
            raise Exception(f"Aba '{nome_aba}' não encontrada. Abas disponíveis: {', '.join(workbook.sheetnames)}")
        
        sheet = workbook[nome_aba]
        valores = []
        for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0] is None:
                break
            valores.append(str(row[0]))
        
        # Validar se tem dados
        if not valores:
            raise Exception(f"Nenhum dado encontrado na coluna A da aba '{nome_aba}' (a partir da linha 2)")
        
        return valores
    except Exception as e:
        raise Exception(f"Erro ao ler arquivo Excel: {e}")

def salvar_resultados_excel(results, arquivo_saida):
    """Salva os resultados em um arquivo Excel (modo append)"""
    try:
        # Verificar se o arquivo já existe
        if os.path.exists(arquivo_saida):
            # Carregar arquivo existente
            workbook = load_workbook(arquivo_saida)
            sheet = workbook.active
            print(f"Arquivo existente carregado. Adicionando {len(results)} novos registros...")
        else:
            # Criar novo arquivo com cabeçalhos
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Resultados"
            sheet.append(['Protocolo', 'Status', 'Mensagem', 'Data/Hora'])
            print(f"Novo arquivo criado: {arquivo_saida}")
        
        # Adicionar novos dados
        for result in results:
            sheet.append([
                result['Protocolo'], 
                result['Status'], 
                result['Mensagem'],
                result.get('DataHora', '')
            ])
        
        workbook.save(arquivo_saida)
        print(f"✓ {len(results)} registros adicionados ao arquivo: {arquivo_saida}")
        
    except Exception as e:
        print(f"Erro ao salvar resultados: {e}")

# ==================== CONFIGURAÇÃO DO DRIVER ====================
def criar_driver():
    """Cria e configura o driver do Firefox"""
    firefox_options = Options()
    firefox_options.add_argument('--width=1920')
    firefox_options.add_argument('--height=1080')
    # firefox_options.add_argument('--headless')  # Descomente para modo headless
    
    # Configurações adicionais para melhor estabilidade
    firefox_options.set_preference("dom.webdriver.enabled", False)
    firefox_options.set_preference("useAutomationExtension", False)
    firefox_options.set_preference("network.http.connection-timeout", 90)
    firefox_options.set_preference("network.http.response.timeout", 90)
    
    driver = webdriver.Firefox(
        service=FirefoxService(GeckoDriverManager().install()), 
        options=firefox_options
    )
    driver.maximize_window()
    driver.set_page_load_timeout(60)
    return driver

def fechar_driver_seguro(driver, logger):
    """Fecha o driver de forma segura"""
    try:
        if driver:
            logger.escrever("Fechando conexão com Firefox...")
            driver.quit()
            time.sleep(3)  # Firefox precisa de mais tempo para fechar completamente
            logger.escrever("Firefox fechado com sucesso")
    except Exception as e:
        logger.escrever(f"Erro ao fechar Firefox: {e}")

def reconectar_driver(email, senha, logger):
    """
    Reconecta ao sistema criando novo driver e fazendo login
    Retorna: (driver, sipcs_handle, siach_handle) ou (None, None, None) se falhar
    """
    try:
        logger.escrever("Iniciando reconexão...")
        
        # Criar novo driver
        driver = criar_driver()
        logger.escrever("Novo driver criado")
        
        # Fazer login
        fazer_login(driver, email, senha, logger)
        sipcs_handle = driver.current_window_handle
        
        # Acessar SIACH
        acessar_siach(driver, logger)
        
        # Mudar para aba SIACH
        abas_handles = driver.window_handles
        siach_handle = [h for h in abas_handles if h != sipcs_handle][0]
        driver.switch_to.window(siach_handle)
        driver.execute_script(f"document.body.style.zoom='{CONFIG['ZOOM_NIVEL']}'")
        
        logger.escrever("Reconexão concluída com sucesso")
        return driver, sipcs_handle, siach_handle
        
    except Exception as e:
        logger.escrever(f"Falha na reconexão: {e}")
        if driver:
            fechar_driver_seguro(driver, logger)
        return None, None, None

# ==================== FUNÇÕES DE NAVEGAÇÃO ====================
def fazer_login(driver, email, senha, logger):
    """Realiza o login no sistema SIPCS"""
    max_tentativas = 3
    for tentativa in range(max_tentativas):
        try:
            logger.escrever(f"Tentativa de login {tentativa + 1}/{max_tentativas}")
            driver.get(CONFIG['URL_SIPCS'])
            
            usuario_input = WebDriverWait(driver, CONFIG['TIMEOUT_PADRAO']).until(
                EC.presence_of_element_located((By.NAME, 'loginForm:username'))
            )
            senha_input = driver.find_element(By.NAME, 'loginForm:password')
            
            usuario_input.clear()
            senha_input.clear()
            usuario_input.send_keys(email)
            senha_input.send_keys(senha)
            
            botao_login = driver.find_element(By.XPATH, '//*[@id="loginForm"]/div/div/input')
            botao_login.click()
            
            time.sleep(3)
            logger.escrever("Login realizado com sucesso")
            return True
            
        except Exception as e:
            logger.escrever(f"Erro na tentativa {tentativa + 1}: {e}")
            if tentativa < max_tentativas - 1:
                logger.escrever("Aguardando 3 segundos antes de tentar novamente...")
                time.sleep(3)
            else:
                raise Exception(f"Falha ao realizar login após {max_tentativas} tentativas: {e}")

def acessar_siach(driver, logger):
    """Acessa o módulo SIACH"""
    try:
        opcao_siach = WebDriverWait(driver, CONFIG['TIMEOUT_PADRAO']).until(
            EC.element_to_be_clickable((By.XPATH, 
                '//*[@id="index"]/fieldset/div[1]/table/tbody/tr[1]/td[2]/a'))
        )
        opcao_siach.click()
        time.sleep(2)
        return True
    except Exception as e:
        logger.escrever(f"Erro ao acessar SIACH: {e}")
        return False

def clicar_com_javascript(driver, element):
    """Clica em um elemento usando JavaScript (mais confiável)"""
    driver.execute_script("arguments[0].click();", element)
    time.sleep(CONFIG['ESPERA_ENTRE_ACOES'])

# ==================== PROCESSAMENTO DE NOME ====================
def extrair_primeiro_nome(nome_completo):
    """
    Extrai e formata o primeiro nome do cliente
    Exemplo: "FRANCISCA DANIELA VIEIRA" -> "Francisca"
    """
    if not nome_completo:
        return ""
    
    partes = nome_completo.strip().split()
    
    if not partes:
        return ""
    
    primeiro_nome = partes[0].capitalize()
    
    return primeiro_nome

def obter_nome_cliente(driver, logger):
    """
    Obtém o nome do cliente do sistema
    Tenta dois XPaths diferentes
    """
    xpaths_nome = [
        '//*[@id="dados_cliente_2"]/div/div[2]/span',
        '/html/body/div[2]/div/div/div[6]/div/div[2]/fieldset[2]/div[2]/div[1]/div/div[2]/span'
    ]
    
    for xpath in xpaths_nome:
        try:
            elemento = WebDriverWait(driver, 5).until(
                EC.presence_of_element_located((By.XPATH, xpath))
            )
            nome_completo = elemento.text.strip()
            
            if nome_completo:
                logger.escrever(f"Nome do cliente encontrado: {nome_completo}")
                return nome_completo
                
        except (TimeoutException, NoSuchElementException):
            continue
    
    logger.escrever("Nome do cliente não encontrado em nenhum dos XPaths")
    return ""

def processar_script_personalizado(script, nome_completo, logger):
    """
    Processa o script substituindo @NomeCliente@ pelo primeiro nome
    """
    if '@NomeCliente@' not in script:
        return script
    
    primeiro_nome = extrair_primeiro_nome(nome_completo)
    
    if not primeiro_nome:
        logger.escrever("Não foi possível extrair o primeiro nome. Usando script sem substituição.")
        return script.replace('@NomeCliente@', '[Nome não encontrado]')
    
    script_personalizado = script.replace('@NomeCliente@', primeiro_nome)
    logger.escrever(f"Script personalizado: @NomeCliente@ -> {primeiro_nome}")
    
    return script_personalizado

# ==================== PROCESSAMENTO DE PROTOCOLO/OCORRÊNCIA ====================
def processar_protocolo(driver, valor, script, usar_personalizacao, tipo_processamento, logger):
    """Processa um único protocolo ou ocorrência"""
    resultado = {
        'Protocolo': valor,
        'Status': 'Erro',
        'Mensagem': '',
        'DataHora': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    }
    
    try:
        # Acessar opção "Atender Ocorrência"
        opcao_atender = WebDriverWait(driver, CONFIG['TIMEOUT_PADRAO']).until(
            EC.element_to_be_clickable((By.XPATH, '//*[@id="menu"]/li[3]/a'))
        )
        time.sleep(0.5)
        clicar_com_javascript(driver, opcao_atender)
        time.sleep(2)
        
        # Definir XPath baseado no tipo (Protocolo ou Ocorrência)
        if tipo_processamento == "protocolo":
            xpath_campo = '/html/body/div[2]/div/div/div[2]/div[1]/form/div[1]/div[1]/div[3]/div/input'
        else:  # ocorrencia
            xpath_campo = '/html/body/div[2]/div/div/div[2]/div[1]/form/div[1]/div[1]/div[2]/div/input'
        
        # Preencher campo de protocolo ou ocorrência
        campo_texto = driver.find_element(By.XPATH, xpath_campo)
        campo_texto.clear()
        campo_texto.send_keys(valor)
        
        # Consultar
        botao_consultar = driver.find_element(By.XPATH, 
            '/html/body/div[2]/div/div/div[2]/div[1]/form/div[2]/a[1]')
        time.sleep(0.5)
        clicar_com_javascript(driver, botao_consultar)
        
        # Obter fase do protocolo
        fase = WebDriverWait(driver, CONFIG['TIMEOUT_PADRAO']).until(
            EC.presence_of_element_located((By.XPATH, 
                '/html/body/div[2]/div/div/div[2]/div[2]/div/div[2]/div/div/div/table/tbody/tr/td/div/div[1]/div/div[2]/p[3]/span'))
        )
        
        try:
            fase_texto = fase.text
        except StaleElementReferenceException:
            fase = driver.find_element(By.XPATH, 
                '/html/body/div[2]/div/div/div[2]/div[2]/div/div[2]/div/div/div/table/tbody/tr/td/div/div[1]/div/div[2]/p[3]/span')
            fase_texto = fase.text
        
        # Processar de acordo com a fase
        if fase_texto in ["ABERTA", "EM ANDAMENTO"]:
            return finalizar_protocolo(driver, valor, script, usar_personalizacao, fase_texto, logger)
        elif fase_texto in ["REABERTA", "FINALIZADA"]:
            resultado['Status'] = 'Sem ação'
            resultado['Mensagem'] = f'{tipo_processamento.capitalize()} com fase {fase_texto}'
            logger.escrever(f'Sem ação: {tipo_processamento.capitalize()} {valor} com fase {fase_texto}')
        else:
            resultado['Status'] = 'Fase desconhecida'
            resultado['Mensagem'] = f'Fase: {fase_texto}'
            logger.escrever(f'Fase desconhecida: {tipo_processamento.capitalize()} {valor} com fase {fase_texto}')
        
        return resultado
        
    except WebDriverException as e:
        # Verificar se é erro de conexão HTTP - PROPAGAR para reconexão
        erro_msg = str(e)
        if any(erro in erro_msg for erro in ['HTTPConnectionPool', 'Read timed out', 'target frame detached', 'Connection refused', 'Reached error page']):
            logger.escrever(f'Erro de conexão HTTP detectado no {tipo_processamento} {valor}')
            raise  # Propaga o erro para o loop principal reconectar
        else:
            # Outros erros WebDriver - registra e continua
            resultado['Mensagem'] = f'WebDriverException: {erro_msg[:100]}'
            logger.escrever(f'WebDriverException no {tipo_processamento} {valor} - Continuando')
            return resultado
        
    except TimeoutException as e:
        resultado['Mensagem'] = f'Timeout: elemento não encontrado'
        logger.escrever(f'TimeoutException no {tipo_processamento} {valor} - Continuando para o próximo')
        return resultado
        
    except NoSuchElementException as e:
        resultado['Mensagem'] = f'Elemento não encontrado na página'
        logger.escrever(f'NoSuchElementException no {tipo_processamento} {valor} - Continuando para o próximo')
        return resultado
        
    except Exception as e:
        # Verificar se a mensagem de erro indica problema de conexão
        erro_msg = str(e)
        if any(erro in erro_msg for erro in ['HTTPConnectionPool', 'Read timed out', 'Connection refused']):
            logger.escrever(f'Erro de conexão detectado no {tipo_processamento} {valor}')
            raise WebDriverException(erro_msg)  # Converte para WebDriverException e propaga
        else:
            resultado['Mensagem'] = f'Erro: {erro_msg[:100]}'
            logger.escrever(f'Erro ao processar {tipo_processamento} {valor}: {e} - Continuando para o próximo')
            return resultado

def finalizar_protocolo(driver, valor, script, usar_personalizacao, fase_texto, logger):
    """Finaliza um protocolo"""
    resultado = {
        'Protocolo': valor,
        'Status': 'Erro',
        'Mensagem': '',
        'DataHora': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    }
    
    try:
        # Clicar na lupa
        botao_lupa = driver.find_element(By.XPATH, '//*[@id="detalhe_ocorrencia"]/span')
        time.sleep(0.5)
        clicar_com_javascript(driver, botao_lupa)
        time.sleep(2)
        
        # Se usar personalização, obter o nome do cliente
        script_final = script
        if usar_personalizacao and '@NomeCliente@' in script:
            nome_cliente = obter_nome_cliente(driver, logger)
            script_final = processar_script_personalizado(script, nome_cliente, logger)
        
        # Clicar em finalizar
        botao_finalizar = WebDriverWait(driver, CONFIG['TIMEOUT_PADRAO']).until(
            EC.element_to_be_clickable((By.XPATH, 
                '//*[@id="content"]/div/div[6]/div/div[3]/form/a[6]'))
        )
        time.sleep(0.5)
        clicar_com_javascript(driver, botao_finalizar)
        time.sleep(2)
        
        # Preencher campos
        campo_justificativa = driver.find_element(By.XPATH, 
            '/html/body/div[2]/div/div/div[2]/div[1]/form/div[1]/fieldset[1]/div/div[2]/div/input')
        campo_justificativa.send_keys('FIM')
        
        campo_resposta = driver.find_element(By.XPATH, 
            '/html/body/div[2]/div/div/div[2]/div[1]/form/div[1]/fieldset[1]/div/div[4]/div/textarea')
        campo_resposta.send_keys(script_final)
        
        campo_tecnica = driver.find_element(By.XPATH, 
            '/html/body/div[2]/div/div/div[2]/div[1]/form/div[1]/fieldset[1]/div/div[7]/div/textarea')
        campo_tecnica.send_keys(script_final)
        
        time.sleep(3)
        
        # Salvar
        botao_salvar = WebDriverWait(driver, CONFIG['TIMEOUT_SALVAR']).until(
            EC.element_to_be_clickable((By.XPATH, 
                "//a[@data-ng-click='tratarOcorrenciaCtrl.form.submit()']"))
        )
        clicar_com_javascript(driver, botao_salvar)
        time.sleep(2)
        
        # Confirmar
        botao_sim = WebDriverWait(driver, CONFIG['TIMEOUT_SALVAR']).until(
            EC.element_to_be_clickable((By.CSS_SELECTOR, "button.btn.btn-default"))
        )
        clicar_com_javascript(driver, botao_sim)
        time.sleep(2)
        
        resultado['Status'] = 'Concluído'
        resultado['Mensagem'] = f'Protocolo finalizado com sucesso'
        logger.escrever(f'Concluído: Protocolo {valor}')
        
    except WebDriverException as e:
        # Verificar se é erro de conexão HTTP - PROPAGAR para reconexão
        erro_msg = str(e)
        if 'HTTPConnectionPool' in erro_msg or 'Read timed out' in erro_msg or 'target frame detached' in erro_msg:
            logger.escrever(f'⚠ Erro de conexão HTTP ao finalizar protocolo {valor}')
            raise  # ✅ PROPAGA o erro
        else:
            resultado['Mensagem'] = f'WebDriverException ao finalizar: {erro_msg[:100]}'
            logger.escrever(f'WebDriverException ao finalizar {valor}: {erro_msg[:150]}')
            try:
                driver.save_screenshot(f'erro_finalizacao_{valor}.png')
            except:
                pass  # Driver pode já estar fechado
            return resultado
        
    except (TimeoutException, NoSuchElementException) as e:
        resultado['Status'] = 'Já finalizado'
        resultado['Mensagem'] = f'Protocolo já finalizado, fase {fase_texto}'
        logger.escrever(f'Protocolo {valor} já finalizado, fase {fase_texto}')
        
    except Exception as e:
        # Verificar se a mensagem de erro indica problema de conexão
        erro_msg = str(e)
        if 'HTTPConnectionPool' in erro_msg or 'Read timed out' in erro_msg:
            logger.escrever(f'⚠ Erro de conexão ao finalizar protocolo {valor}')
            raise WebDriverException(erro_msg)  # ✅ Converte e propaga
        else:
            resultado['Mensagem'] = f'Erro ao finalizar: {erro_msg[:100]}'
            logger.escrever(f'Erro ao finalizar protocolo {valor}: {e}')
            try:
                driver.save_screenshot(f'erro_finalizacao_{valor}.png')
            except:
                pass  # Driver pode já estar fechado
    
    return resultado

# ==================== FUNÇÃO PRINCIPAL ====================
def fechamento_em_lote(caminho_arquivo, script, email, senha, usar_personalizacao, tipo_processamento,
                      atualizar_callback, progress_callback, stats_callback):
    """Função principal de processamento em lote"""
    
    results = []
    protocolos_processados = 0
    driver = None
    tentativas_reconexao = 0
    
    # Resetar estatísticas no início
    estatisticas.resetar()
    stats_callback()
    
    with GerenciadorLog(CONFIG['ARQUIVO_LOG']) as logger:
        try:
            logger.escrever("=== Iniciando processamento em lote ===")
            logger.escrever(f"Tipo de processamento: {tipo_processamento.upper()}")
            logger.escrever(f"Personalização de nome: {'ATIVADA' if usar_personalizacao else 'DESATIVADA'}")
            
            # Ler protocolos do Excel
            valores_coluna_a = ler_valores_coluna_a(caminho_arquivo, 
                                                    CONFIG['NOME_ABA_EXCEL'])
            total_protocolos = len(valores_coluna_a)
            tipo_texto = "protocolos" if tipo_processamento == "protocolo" else "ocorrências"
            logger.escrever(f"Total de {tipo_texto} a processar: {total_protocolos}")
            
            # Criar driver
            driver = criar_driver()
            
            # Fazer login
            fazer_login(driver, email, senha, logger)
            sipcs_handle = driver.current_window_handle
            
            # Acessar SIACH
            acessar_siach(driver, logger)
            
            # Mudar para aba SIACH
            abas_handles = driver.window_handles
            siach_handle = [h for h in abas_handles if h != sipcs_handle][0]
            driver.switch_to.window(siach_handle)
            driver.execute_script(f"document.body.style.zoom='{CONFIG['ZOOM_NIVEL']}'")
            
            # Carregar progresso
            progresso = carregar_progresso()
            tipo_texto_singular = "protocolo" if tipo_processamento == "protocolo" else "ocorrência"
            logger.escrever(f"Retomando do {tipo_texto_singular} {progresso + 1}")
            
            # Processar protocolos
            inicio = time.time()
            
            i = progresso
            while i < total_protocolos:
                if not estado_app.esta_executando():
                    logger.escrever("Processamento interrompido pelo usuário")
                    break
                
                valor = valores_coluna_a[i]
                tipo_texto_singular = "protocolo" if tipo_processamento == "protocolo" else "ocorrência"
                logger.escrever(f"Processando {tipo_texto_singular} {i + 1}/{total_protocolos}: {valor}")
                
                try:
                    resultado = processar_protocolo(driver, valor, script, usar_personalizacao, tipo_processamento, logger)
                    results.append(resultado)
                    
                    # Atualizar estatísticas
                    estatisticas.adicionar(resultado['Status'])
                    stats_callback()
                    
                    tipo_texto_singular = "Protocolo" if tipo_processamento == "protocolo" else "Ocorrência"
                    atualizar_callback(f"{resultado['Status']}: {tipo_texto_singular} {valor}")
                    
                    protocolos_processados += 1
                    
                    # ✅ SALVAR PROGRESSO para QUALQUER status (Concluído, Sem ação, Erro, etc)
                    # Isso garante que não ficamos presos no mesmo protocolo com erro
                    # EXCETO em caso de erro de conexão, que vai para o except abaixo
                    salvar_progresso(i + 1)
                    progress_callback((i + 1) / total_protocolos, i + 1, total_protocolos)
                    
                    # Resetar contador de reconexão em caso de sucesso
                    tentativas_reconexao = 0
                    
                    # Avançar para o próximo protocolo
                    i += 1
                    
                except WebDriverException as e:
                    # Erro de conexão HTTP com Firefox - tentar reconectar
                    erro_msg = str(e)
                    # Lista completa de erros de conexão conhecidos
                    erros_conexao = ['HTTPConnectionPool', 'Read timed out', 'target frame detached', 
                                   'Connection refused', 'Reached error page', 'connection was refused']
                    
                    if any(erro in erro_msg for erro in erros_conexao):
                        logger.escrever(f"ERRO DE CONEXAO detectado: {erro_msg[:150]}")
                        
                        if tentativas_reconexao < CONFIG['MAX_TENTATIVAS_RECONEXAO']:
                            tentativas_reconexao += 1
                            logger.escrever(f"Tentativa de reconexão {tentativas_reconexao}/{CONFIG['MAX_TENTATIVAS_RECONEXAO']}")
                            
                            # Fechar driver atual
                            fechar_driver_seguro(driver, logger)
                            
                            # Aguardar antes de reconectar (verificando se usuário não parou)
                            logger.escrever("Aguardando 5 segundos antes de reconectar...")
                            for segundo in range(5):
                                if not estado_app.esta_executando():
                                    logger.escrever("Reconexão cancelada pelo usuário")
                                    return
                                time.sleep(1)
                            
                            # Tentar reconectar
                            driver_novo, sipcs_novo, siach_novo = reconectar_driver(email, senha, logger)
                            
                            if driver_novo:
                                # ATUALIZAR as variáveis com os novos handles
                                driver = driver_novo
                                sipcs_handle = sipcs_novo
                                siach_handle = siach_novo
                                
                                tipo_texto_singular = "protocolo" if tipo_processamento == "protocolo" else "ocorrência"
                                atualizar_callback(f"Reconectado! Retomando {tipo_texto_singular} {valor}")
                                # NÃO incrementa i - vai tentar o mesmo protocolo novamente
                                continue
                            else:
                                logger.escrever("Falha na reconexão. Encerrando processamento.")
                                break
                        else:
                            logger.escrever(f"Número máximo de tentativas de reconexão atingido ({CONFIG['MAX_TENTATIVAS_RECONEXAO']})")
                            resultado = {
                                'Protocolo': valor,
                                'Status': 'Erro',
                                'Mensagem': 'Falha de conexão - max tentativas de reconexão atingidas',
                                'DataHora': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                            }
                            results.append(resultado)
                            estatisticas.adicionar('Erro')
                            stats_callback()
                            salvar_progresso(i + 1)  # Salvar progresso para pular este protocolo problemático
                            i += 1  # Avançar para o próximo
                            tentativas_reconexao = 0  # Resetar para próximo protocolo
                    else:
                        # Outro tipo de WebDriverException
                        logger.escrever(f"WebDriverException: {erro_msg[:150]}")
                        resultado = {
                            'Protocolo': valor,
                            'Status': 'Erro',
                            'Mensagem': f'WebDriverException: {erro_msg[:100]}',
                            'DataHora': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                        }
                        results.append(resultado)
                        estatisticas.adicionar('Erro')
                        stats_callback()
                        salvar_progresso(i + 1)
                        i += 1  # Avança para o próximo
                
            # Finalização
            fim = time.time()
            tempo_total = fim - inicio
            
            # Obter estatísticas finais
            totais = estatisticas.obter_totais()
            
            tipo_texto = "protocolos" if tipo_processamento == "protocolo" else "ocorrências"
            logger.escrever(f"Processamento finalizado em {tempo_total:.2f} segundos")
            logger.escrever(f"{tipo_texto.capitalize()} processados: {protocolos_processados}")
            logger.escrever("=== ESTATÍSTICAS FINAIS ===")
            logger.escrever(f"✓ Concluídos: {totais['concluidos']}")
            logger.escrever(f"○ Sem ação: {totais['sem_acao']}")
            logger.escrever(f"? Fase desconhecida: {totais['fase_desconhecida']}")
            logger.escrever(f"✗ Erros: {totais['erros']}")
            logger.escrever(f"TOTAL: {totais['total']}")
            
            # Salvar resultados
            salvar_resultados_excel(results, CONFIG['ARQUIVO_RESULTADO'])
            
            messagebox.showinfo("Concluído", 
                f"Processamento concluído!\n\n"
                f"✓ Concluídos: {totais['concluidos']}\n"
                f"○ Sem ação: {totais['sem_acao']}\n"
                f"? Fase desconhecida: {totais['fase_desconhecida']}\n"
                f"✗ Erros: {totais['erros']}\n"
                f"━━━━━━━━━━━━━━━\n"
                f"Total: {totais['total']} {tipo_texto}")
            
        except Exception as e:
            logger.escrever(f"Erro fatal: {e}")
            messagebox.showerror("Erro", f"Erro durante o processamento:\n{str(e)[:200]}")
            
        finally:
            if driver:
                fechar_driver_seguro(driver, logger)

# ==================== INTERFACE GRÁFICA ====================
def criar_interface():
    """Cria a interface gráfica"""
    
    ctk.set_appearance_mode("dark")
    ctk.set_default_color_theme("blue")
    
    root = ctk.CTk()
    root.title("WallBot - Finalização de Protocolos/Ocorrências SIACH v2.1 (Firefox)")
    root.geometry("700x620")
    
    # Definir ícone da janela
    try:
        root.iconbitmap("icon.ico")
    except Exception as e:
        print(f"Não foi possível carregar o ícone: {e}")
        # Continua sem ícone se houver erro
    
    # Frame principal superior com 2 colunas
    frame_main = ctk.CTkFrame(root)
    frame_main.pack(padx=20, pady=10, fill="both", expand=False)
    
    # ========== COLUNA ESQUERDA: Formulário ==========
    frame_form = ctk.CTkFrame(frame_main)
    frame_form.grid(row=0, column=0, padx=(0, 10), pady=0, sticky="nsew")
    
    # Email
    ctk.CTkLabel(frame_form, text="Email:", width=90, anchor="w").grid(row=0, column=0, sticky="w", pady=8, padx=10)
    entrada_email = ctk.CTkEntry(frame_form, width=280)
    entrada_email.grid(row=0, column=1, pady=8, padx=(0, 10), sticky="ew")
    
    # Senha
    ctk.CTkLabel(frame_form, text="Senha:", width=90, anchor="w").grid(row=1, column=0, sticky="w", pady=8, padx=10)
    entrada_senha = ctk.CTkEntry(frame_form, width=280, show="*")
    entrada_senha.grid(row=1, column=1, pady=8, padx=(0, 10), sticky="ew")
    
    # Arquivo Excel
    ctk.CTkLabel(frame_form, text="Arquivo Excel:", width=90, anchor="w").grid(row=2, column=0, sticky="w", pady=8, padx=10)
    
    frame_arquivo = ctk.CTkFrame(frame_form, fg_color="transparent")
    frame_arquivo.grid(row=2, column=1, pady=8, padx=(0, 10), sticky="ew")
    
    entrada_arquivo = ctk.CTkEntry(frame_arquivo)
    entrada_arquivo.pack(side="left", fill="x", expand=True, padx=(0, 5))
    
    def selecionar_arquivo():
        arquivo = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls")])
        if arquivo:
            entrada_arquivo.delete(0, "end")
            entrada_arquivo.insert(0, arquivo)
    
    ctk.CTkButton(
        frame_arquivo,
        text="Selecionar",
        width=80,
        command=selecionar_arquivo
    ).pack(side="left")
    
    # Tipo de Processamento (Protocolo ou Ocorrência)
    ctk.CTkLabel(frame_form, text="Processar:", width=90, anchor="w").grid(row=3, column=0, sticky="w", pady=8, padx=10)
    
    tipo_proc_var = ctk.StringVar(value="protocolo")
    
    frame_tipo_proc = ctk.CTkFrame(frame_form, fg_color="transparent")
    frame_tipo_proc.grid(row=3, column=1, sticky="w", padx=(0, 10), pady=8)
    
    radio_protocolo = ctk.CTkRadioButton(
        frame_tipo_proc,
        text="Protocolos",
        variable=tipo_proc_var,
        value="protocolo"
    )
    radio_protocolo.pack(side="left", padx=(0, 15))
    
    radio_ocorrencia = ctk.CTkRadioButton(
        frame_tipo_proc,
        text="Ocorrências",
        variable=tipo_proc_var,
        value="ocorrencia"
    )
    radio_ocorrencia.pack(side="left")
    
    # Script
    ctk.CTkLabel(frame_form, text="Script:", width=90, anchor="w").grid(row=4, column=0, sticky="nw", pady=8, padx=10)
    
    frame_script = ctk.CTkFrame(frame_form, fg_color="transparent")
    frame_script.grid(row=4, column=1, pady=8, padx=(0, 10), sticky="ew")
    
    entrada_script = ctk.CTkTextbox(frame_script, height=90)
    entrada_script.pack(side="left", fill="both", expand=True, padx=(0, 5))
    
    ctk.CTkButton(
        frame_script,
        text="Limpar",
        width=80,
        command=lambda: entrada_script.delete("1.0", "end")
    ).pack(side="left")
    
    # Radio Button para tipo de script
    ctk.CTkLabel(frame_form, text="Tipo:", width=90, anchor="w").grid(row=5, column=0, sticky="nw", pady=8, padx=10)
    
    tipo_script_var = ctk.StringVar(value="padrao")
    
    frame_radio = ctk.CTkFrame(frame_form, fg_color="transparent")
    frame_radio.grid(row=5, column=1, sticky="w", padx=(0, 10), pady=8)
    
    radio_padrao = ctk.CTkRadioButton(
        frame_radio,
        text="Padrão",
        variable=tipo_script_var,
        value="padrao"
    )
    radio_padrao.pack(anchor="w", pady=2)
    
    radio_personalizado = ctk.CTkRadioButton(
        frame_radio,
        text="Personalizado (@NomeCliente@)",
        variable=tipo_script_var,
        value="personalizado"
    )
    radio_personalizado.pack(anchor="w", pady=2)
    
    # Label explicativa
    label_explicacao = ctk.CTkLabel(
        frame_form,
        text="💡 Ex: 'Olá @NomeCliente@, resolvido!'",
        font=("Arial", 9),
        text_color="gray"
    )
    label_explicacao.grid(row=6, column=0, columnspan=2, sticky="w", padx=10, pady=(0, 5))
    
    # Configurar expansão da coluna de entrada
    frame_form.columnconfigure(1, weight=1)
    
    # ========== COLUNA DIREITA: Estatísticas ==========
    frame_stats_container = ctk.CTkFrame(frame_main)
    frame_stats_container.grid(row=0, column=1, padx=(10, 0), pady=0, sticky="nsew")
    
    ctk.CTkLabel(
        frame_stats_container, 
        text="📊 ESTATÍSTICAS",
        font=("Arial", 13, "bold")
    ).pack(pady=(10, 15))
    
    # Concluídos
    frame_concluidos = ctk.CTkFrame(frame_stats_container)
    frame_concluidos.pack(fill="x", padx=12, pady=5)
    ctk.CTkLabel(frame_concluidos, text="✓ Concluídos:", font=("Arial", 11, "bold"), width=120, anchor="w").pack(side="left", padx=5)
    label_concluidos = ctk.CTkLabel(frame_concluidos, text="0", font=("Arial", 11), text_color="#00FF00", width=40, anchor="e")
    label_concluidos.pack(side="right", padx=5)
    
    # Sem ação
    frame_sem_acao = ctk.CTkFrame(frame_stats_container)
    frame_sem_acao.pack(fill="x", padx=12, pady=5)
    ctk.CTkLabel(frame_sem_acao, text="○ Sem ação:", font=("Arial", 11, "bold"), width=120, anchor="w").pack(side="left", padx=5)
    label_sem_acao = ctk.CTkLabel(frame_sem_acao, text="0", font=("Arial", 11), text_color="#FFD700", width=40, anchor="e")
    label_sem_acao.pack(side="right", padx=5)
    
    # Fase desconhecida
    frame_fase_desc = ctk.CTkFrame(frame_stats_container)
    frame_fase_desc.pack(fill="x", padx=12, pady=5)
    ctk.CTkLabel(frame_fase_desc, text="? Fase desconhecida:", font=("Arial", 11, "bold"), width=120, anchor="w").pack(side="left", padx=5)
    label_fase_desc = ctk.CTkLabel(frame_fase_desc, text="0", font=("Arial", 11), text_color="#FFA500", width=40, anchor="e")
    label_fase_desc.pack(side="right", padx=5)
    
    # Erros
    frame_erros = ctk.CTkFrame(frame_stats_container)
    frame_erros.pack(fill="x", padx=12, pady=5)
    ctk.CTkLabel(frame_erros, text="✗ Erros:", font=("Arial", 11, "bold"), width=120, anchor="w").pack(side="left", padx=5)
    label_erros = ctk.CTkLabel(frame_erros, text="0", font=("Arial", 11), text_color="#FF4444", width=40, anchor="e")
    label_erros.pack(side="right", padx=5)
    
    # Separador
    ctk.CTkLabel(frame_stats_container, text="━━━━━━━━━━━━━━", font=("Arial", 10)).pack(pady=8)
    
    # Total
    frame_total = ctk.CTkFrame(frame_stats_container)
    frame_total.pack(fill="x", padx=12, pady=5)
    ctk.CTkLabel(frame_total, text="TOTAL:", font=("Arial", 12, "bold"), width=120, anchor="w").pack(side="left", padx=5)
    label_total = ctk.CTkLabel(frame_total, text="0", font=("Arial", 12, "bold"), text_color="#FFFFFF", width=40, anchor="e")
    label_total.pack(side="right", padx=5)
    
    # Configurar peso das colunas - formulário expande mais
    frame_main.columnconfigure(0, weight=2)
    frame_main.columnconfigure(1, weight=1)
    
    # Frame para os botões
    frame_botoes = ctk.CTkFrame(root)
    frame_botoes.pack(padx=20, pady=10, fill="x")
    
    def iniciar_processamento():
        email = entrada_email.get().strip()
        senha = entrada_senha.get().strip()
        arquivo = entrada_arquivo.get().strip()
        script = entrada_script.get("1.0", "end-1c").strip()
        usar_personalizacao = tipo_script_var.get() == "personalizado"
        tipo_processamento = tipo_proc_var.get()
        
        # Validações
        if not email or not senha:
            messagebox.showwarning("Aviso", "Por favor, preencha email e senha!")
            return
        
        # Validação básica de email
        if '@' not in email or '.' not in email:
            messagebox.showwarning("Aviso", "Email parece inválido!")
            return
            
        if not arquivo:
            messagebox.showwarning("Aviso", "Por favor, selecione um arquivo Excel!")
            return
        
        # Validar se arquivo existe
        if not os.path.exists(arquivo):
            messagebox.showerror("Erro", f"Arquivo não encontrado:\n{arquivo}")
            return
        
        # Validar extensão do arquivo
        if not arquivo.lower().endswith(('.xlsx', '.xls')):
            messagebox.showwarning("Aviso", "Por favor, selecione um arquivo Excel válido (.xlsx ou .xls)!")
            return
            
        if not script:
            messagebox.showwarning("Aviso", "Por favor, preencha o script!")
            return
        
        # Validar se o script personalizado contém @NomeCliente@
        if usar_personalizacao and '@NomeCliente@' not in script:
            resposta = messagebox.askyesno(
                "Script Personalizado",
                "Você selecionou 'Personalizado' mas o script não contém @NomeCliente@.\n\n"
                "Deseja continuar mesmo assim?"
            )
            if not resposta:
                return
        
        estado_app.iniciar()
        threading.Thread(target=fechamento_em_lote, args=(
            arquivo, script, email, senha, usar_personalizacao, tipo_processamento,
            atualizar_resultado, atualizar_progresso, atualizar_estatisticas
        ), daemon=True).start()
    
    def parar_processamento():
        estado_app.parar()
        atualizar_resultado("Parando processamento...")
    
    def resetar():
        resetar_progresso()
        estatisticas.resetar()
        progress_bar.set(0)
        progress_label.configure(text="Progresso: 0/0")
        text_resultado.delete("1.0", "end")
        atualizar_estatisticas()
        atualizar_resultado("Progresso e estatísticas resetados!")
    
    def limpar_historico():
        if messagebox.askyesno("Confirmar", 
            "Deseja realmente apagar o arquivo de resultados?\nEsta ação não pode ser desfeita!"):
            try:
                if os.path.exists(CONFIG['ARQUIVO_RESULTADO']):
                    os.remove(CONFIG['ARQUIVO_RESULTADO'])
                    atualizar_resultado(f"✓ Arquivo {CONFIG['ARQUIVO_RESULTADO']} removido com sucesso!")
                else:
                    atualizar_resultado(f"Arquivo {CONFIG['ARQUIVO_RESULTADO']} não existe.")
            except Exception as e:
                messagebox.showerror("Erro", f"Erro ao remover arquivo: {e}")
    
    def atualizar_resultado(texto):
        text_resultado.insert("end", texto + "\n")
        text_resultado.see("end")
    
    def atualizar_progresso(percentual, atual, total):
        progress_bar.set(percentual)
        progress_label.configure(text=f"Progresso: {atual}/{total}")
    
    def atualizar_estatisticas():
        totais = estatisticas.obter_totais()
        label_concluidos.configure(text=str(totais['concluidos']))
        label_sem_acao.configure(text=str(totais['sem_acao']))
        label_fase_desc.configure(text=str(totais['fase_desconhecida']))
        label_erros.configure(text=str(totais['erros']))
        label_total.configure(text=str(totais['total']))
    
    ctk.CTkButton(
        frame_botoes, text="Iniciar", 
        command=iniciar_processamento, width=120
    ).pack(side="left", padx=5)
    
    ctk.CTkButton(
        frame_botoes, text="Parar", 
        command=parar_processamento, width=120
    ).pack(side="left", padx=5)
    
    ctk.CTkButton(
        frame_botoes, text="Resetar Progresso", 
        command=resetar, width=140
    ).pack(side="left", padx=5)
    
    ctk.CTkButton(
        frame_botoes, text="Limpar Histórico", 
        command=limpar_historico, width=140
    ).pack(side="left", padx=5)
    
    # Barra de progresso e rótulo
    frame_progresso = ctk.CTkFrame(root)
    frame_progresso.pack(padx=20, pady=10, fill="x")

    progress_label = ctk.CTkLabel(frame_progresso, text="Progresso: 0/0")
    progress_label.pack(anchor="w")

    progress_bar = ctk.CTkProgressBar(frame_progresso, width=400)
    progress_bar.pack(pady=5, fill="x")
    progress_bar.set(0)
    
    # Resultado
    frame_resultado = ctk.CTkFrame(root)
    frame_resultado.pack(padx=20, pady=10, fill="both", expand=True)
    
    ctk.CTkLabel(frame_resultado, text="📋 Log de Processamento", font=("Arial", 12, "bold")).pack(anchor="w", padx=10, pady=5)
    
    text_resultado = ctk.CTkTextbox(frame_resultado, wrap="word")
    text_resultado.pack(fill="both", expand=True, padx=10, pady=5)
    
    root.mainloop()

# ==================== EXECUÇÃO ====================
if __name__ == "__main__":
    criar_interface()